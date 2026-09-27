"""Validate the reviewed E2E pilot workbook and freeze it into JSONL."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


CATEGORIES = {
    "stable_fact", "long_term_preference", "long_term_goal",
    "temporary_state", "one_off_information", "small_talk",
}
ENTITY_TYPES = {
    "person", "project", "technology", "tool", "preference", "goal",
    "topic", "organization", "place", "event", "other",
}
PREDICATES = {
    "prefers", "dislikes", "uses", "learns", "builds", "works_on",
    "plans", "knows", "experienced", "belongs_to", "related_to",
}
IMPORTANCE_VALUES = {0.1, 0.3, 0.5, 0.7, 0.9}
SOURCE_FIELDS = (
    "sample_id", "user_message", "assistant_context", "category", "difficulty",
)
HEADERS = SOURCE_FIELDS + (
    "human_should_remember", "human_importance", "expected_entities",
    "expected_relations", "annotation_reason", "review_status",
)


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    if any(not line.strip() for line in lines):
        raise ValueError("blank JSONL lines are not allowed")
    return [json.loads(line) for line in lines]


def parse_json_array(value: Any, sample_id: str, field: str) -> list[dict[str, Any]]:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{sample_id}: {field} must contain JSON")
    try:
        parsed = json.loads(value)
    except json.JSONDecodeError as exc:
        raise ValueError(f"{sample_id}: invalid {field} JSON: {exc}") from exc
    if not isinstance(parsed, list):
        raise ValueError(f"{sample_id}: {field} must be a JSON array")
    return parsed


def validate_entity_relation_fields(row: dict[str, Any]) -> None:
    sample_id = row["sample_id"]
    entities = row["expected_entities"]
    relations = row["expected_relations"]
    if len(entities) > 8:
        raise ValueError(f"{sample_id}: expected_entities exceeds production limit 8")
    if len(relations) > 10:
        raise ValueError(f"{sample_id}: expected_relations exceeds production limit 10")

    keys: set[str] = set()
    for entity in entities:
        if not isinstance(entity, dict):
            raise ValueError(f"{sample_id}: expected entity must be an object")
        key = entity.get("entity_key")
        if not isinstance(key, str) or not key or key in keys:
            raise ValueError(f"{sample_id}: invalid or duplicate entity_key")
        if entity.get("type") not in ENTITY_TYPES:
            raise ValueError(f"{sample_id}: invalid entity type {entity.get('type')!r}")
        if not isinstance(entity.get("name"), str) or not entity["name"].strip():
            raise ValueError(f"{sample_id}: entity name is missing")
        keys.add(key)

    for relation in relations:
        if not isinstance(relation, dict):
            raise ValueError(f"{sample_id}: expected relation must be an object")
        if relation.get("predicate") not in PREDICATES:
            raise ValueError(f"{sample_id}: invalid predicate {relation.get('predicate')!r}")
        source = relation.get("source_entity_key")
        target = relation.get("target_entity_key")
        if source != "user" and source not in keys:
            raise ValueError(f"{sample_id}: invalid relation source {source!r}")
        if target != "user" and target not in keys:
            raise ValueError(f"{sample_id}: invalid relation target {target!r}")


def read_and_validate_workbook(workbook_path: Path, source_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    if "Annotation" not in workbook.sheetnames:
        raise ValueError("workbook is missing Annotation sheet")
    sheet = workbook["Annotation"]
    values = list(sheet.iter_rows(values_only=True))
    if not values or tuple(values[0]) != HEADERS:
        raise ValueError("Annotation headers do not match the required schema")
    if len(values) - 1 != 30 or len(source_rows) != 30:
        raise ValueError("workbook and source JSONL must each contain 30 samples")

    expected_ids = [f"p{index:03d}" for index in range(1, 31)]
    frozen: list[dict[str, Any]] = []
    for index, excel_values in enumerate(values[1:]):
        excel = dict(zip(HEADERS, excel_values, strict=True))
        source = source_rows[index]
        sample_id = excel["sample_id"]
        if sample_id != expected_ids[index] or source.get("sample_id") != sample_id:
            raise ValueError("sample IDs must be exactly p001..p030 in order")
        for field in SOURCE_FIELDS:
            if excel[field] != source.get(field):
                raise ValueError(f"{sample_id}: original field changed: {field}")

        if type(excel["human_should_remember"]) is not bool:
            raise ValueError(f"{sample_id}: human_should_remember must be boolean")
        importance = excel["human_importance"]
        if not isinstance(importance, (int, float)) or float(importance) not in IMPORTANCE_VALUES:
            raise ValueError(f"{sample_id}: invalid human_importance")
        if excel["review_status"] != "confirmed":
            raise ValueError(f"{sample_id}: review_status is not confirmed")
        reason = excel["annotation_reason"]
        if not isinstance(reason, str) or not reason.strip():
            raise ValueError(f"{sample_id}: annotation_reason is missing")

        row = {field: source[field] for field in SOURCE_FIELDS}
        row.update({
            "human_should_remember": excel["human_should_remember"],
            "human_importance": float(importance),
            "expected_entities": parse_json_array(excel["expected_entities"], sample_id, "expected_entities"),
            "expected_relations": parse_json_array(excel["expected_relations"], sample_id, "expected_relations"),
            "annotation_reason": reason,
            "review_status": "confirmed",
        })
        validate_entity_relation_fields(row)
        frozen.append(row)

    categories = Counter(row["category"] for row in frozen)
    if set(categories) != CATEGORIES or any(count != 5 for count in categories.values()):
        raise ValueError("each of the six categories must contain exactly five samples")
    return frozen


def render_summary(rows: list[dict[str, Any]], workbook_path: Path, dataset_path: Path) -> str:
    status = Counter(row["review_status"] for row in rows)
    should = Counter(row["human_should_remember"] for row in rows)
    importance = Counter(row["human_importance"] for row in rows)
    categories = Counter(row["category"] for row in rows)
    entity_count = sum(len(row["expected_entities"]) for row in rows)
    relation_count = sum(len(row["expected_relations"]) for row in rows)
    false_kg = [
        row["sample_id"] for row in rows
        if not row["human_should_remember"]
        and (row["expected_entities"] or row["expected_relations"])
    ]
    lines = [
        "# E2E Memory Pilot Ground Truth Summary",
        "",
        "## Freeze status",
        "",
        f"- Source workbook: `{workbook_path.as_posix()}`",
        f"- Frozen dataset: `{dataset_path.as_posix()}`",
        f"- Samples: {len(rows)}",
        f"- Confirmed: {status.get('confirmed', 0)}",
        f"- Pending: {status.get('pending_human_annotation', 0)}",
        "- Ground Truth status: confirmed",
        "- AI scoring: not run",
        "- Database write: not performed",
        "",
        "## Label distributions",
        "",
        f"- shouldRemember=true: {should.get(True, 0)}",
        f"- shouldRemember=false: {should.get(False, 0)}",
        "- Importance: " + ", ".join(f"{value:.1f}={importance.get(value, 0)}" for value in sorted(IMPORTANCE_VALUES)),
        "",
        "## Category distribution",
        "",
    ]
    lines.extend(f"- {category}: {categories[category]}" for category in sorted(categories))
    lines.extend([
        "",
        "## Knowledge graph Ground Truth",
        "",
        f"- Expected entity total: {entity_count}",
        f"- Expected relation total: {relation_count}",
        "- ShouldRemember=false with non-empty Expected KG: " + ", ".join(false_kg),
        "",
        "The E2E Pilot Ground Truth was produced and confirmed by one human annotator.",
        "",
    ])
    return "\n".join(lines)


def write_atomically(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(content, encoding="utf-8")
    temporary.replace(path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--dataset", required=True, type=Path)
    parser.add_argument("--report", required=True, type=Path)
    args = parser.parse_args()

    source_rows = load_jsonl(args.dataset)
    frozen = read_and_validate_workbook(args.workbook, source_rows)
    dataset_text = "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in frozen)
    report_text = render_summary(frozen, args.workbook, args.dataset)
    write_atomically(args.dataset, dataset_text)
    write_atomically(args.report, report_text)

    result = {
        "samples": len(frozen),
        "confirmed": sum(row["review_status"] == "confirmed" for row in frozen),
        "pending": sum(row["review_status"] != "confirmed" for row in frozen),
        "shouldRemember": dict(Counter(str(row["human_should_remember"]).lower() for row in frozen)),
        "importance": {f"{value:.1f}": sum(row["human_importance"] == value for row in frozen) for value in sorted(IMPORTANCE_VALUES)},
        "entities": sum(len(row["expected_entities"]) for row in frozen),
        "relations": sum(len(row["expected_relations"]) for row in frozen),
        "false_with_kg": [row["sample_id"] for row in frozen if not row["human_should_remember"] and (row["expected_entities"] or row["expected_relations"])],
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
